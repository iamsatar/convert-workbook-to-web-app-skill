#!/usr/bin/env python3
"""Create a structural and formula inventory for an OOXML Excel workbook."""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any


SUPPORTED_SUFFIXES = {".xlsx", ".xlsm", ".xltx", ".xltm"}
FUNCTION_RE = re.compile(r"\b([A-Z][A-Z0-9._]*)\s*\(", re.IGNORECASE)
CELL_REF_RE = re.compile(
    r"(?:(?:'([^']+)'|([A-Za-z_][A-Za-z0-9_ .]*))!)?(\$?[A-Z]{1,3}\$?\d+)",
    re.IGNORECASE,
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Inspect workbook structure, formulas, validations, and OOXML features."
    )
    parser.add_argument("workbook", type=Path, help="Path to an .xlsx or .xlsm workbook")
    parser.add_argument("--output", type=Path, help="Write JSON to this file instead of stdout")
    parser.add_argument(
        "--include-values",
        action="store_true",
        help="Include non-formula cell values; may expose workbook data and create a large report",
    )
    parser.add_argument(
        "--label-limit",
        type=int,
        default=2000,
        help="Maximum number of text labels retained across the workbook (default: 2000)",
    )
    return parser.parse_args()


def safe_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def inspect_package(path: Path) -> dict[str, Any]:
    feature_prefixes = {
        "pivot_tables": "xl/pivotTables/",
        "pivot_caches": "xl/pivotCache/",
        "external_links": "xl/externalLinks/",
        "query_tables": "xl/queryTables/",
        "slicers": "xl/slicers/",
    }
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
    return {
        "has_vba": "xl/vbaProject.bin" in names,
        "has_connections": "xl/connections.xml" in names,
        "has_calculation_chain": "xl/calcChain.xml" in names,
        **{
            feature: sum(name.startswith(prefix) for name in names)
            for feature, prefix in feature_prefixes.items()
        },
    }


def defined_names(wb: Any) -> list[dict[str, Any]]:
    collection = wb.defined_names
    if hasattr(collection, "values"):
        items = list(collection.values())
    else:
        items = list(getattr(collection, "definedName", []))
    result = []
    for item in items:
        result.append(
            {
                "name": getattr(item, "name", None),
                "value": getattr(item, "attr_text", None) or getattr(item, "value", None),
                "scope_sheet_index": getattr(item, "localSheetId", None),
                "hidden": bool(getattr(item, "hidden", False)),
            }
        )
    return result


def formula_dependencies(formula: str, current_sheet: str) -> list[str]:
    dependencies = set()
    for quoted_sheet, bare_sheet, cell in CELL_REF_RE.findall(formula):
        sheet = quoted_sheet or bare_sheet or current_sheet
        dependencies.add(f"{sheet}!{cell.upper()}")
    return sorted(dependencies)


def inspect_workbook(path: Path, include_values: bool, label_limit: int) -> dict[str, Any]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl is required: install it in an isolated environment") from exc

    keep_vba = path.suffix.lower() in {".xlsm", ".xltm"}
    wb = openpyxl.load_workbook(path, data_only=False, keep_vba=keep_vba, keep_links=True)
    cached_wb = openpyxl.load_workbook(path, data_only=True, read_only=False, keep_links=True)

    formulas: list[dict[str, Any]] = []
    labels: list[dict[str, str]] = []
    values: list[dict[str, Any]] = []
    function_counts: Counter[str] = Counter()
    sheet_reports: list[dict[str, Any]] = []

    for ws in wb.worksheets:
        cached_ws = cached_wb[ws.title]
        sheet_formula_count = 0
        nonempty_count = 0
        comment_count = 0
        input_count = 0
        unlocked_count = 0

        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                nonempty_count += 1
                if cell.comment is not None:
                    comment_count += 1
                if cell.protection.locked is False:
                    unlocked_count += 1
                if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    sheet_formula_count += 1
                    formula_text = str(cell.value)
                    for function in FUNCTION_RE.findall(formula_text):
                        function_counts[function.upper()] += 1
                    formulas.append(
                        {
                            "sheet": ws.title,
                            "cell": cell.coordinate,
                            "formula": formula_text,
                            "cached_value": safe_value(cached_ws[cell.coordinate].value),
                            "number_format": cell.number_format,
                            "dependencies": formula_dependencies(formula_text, ws.title),
                        }
                    )
                else:
                    input_count += 1
                    if isinstance(cell.value, str) and len(labels) < label_limit:
                        labels.append(
                            {
                                "sheet": ws.title,
                                "cell": cell.coordinate,
                                "text": cell.value[:500],
                            }
                        )
                    if include_values:
                        values.append(
                            {
                                "sheet": ws.title,
                                "cell": cell.coordinate,
                                "value": safe_value(cell.value),
                                "number_format": cell.number_format,
                            }
                        )

        validations = []
        data_validations = getattr(ws, "data_validations", None)
        for validation in getattr(data_validations, "dataValidation", []) or []:
            validations.append(
                {
                    "type": validation.type,
                    "range": str(validation.sqref),
                    "formula1": validation.formula1,
                    "formula2": validation.formula2,
                    "allow_blank": validation.allowBlank,
                }
            )

        tables = []
        for table in ws.tables.values():
            tables.append({"name": table.name, "display_name": table.displayName, "range": table.ref})

        sheet_reports.append(
            {
                "name": ws.title,
                "state": ws.sheet_state,
                "used_range": ws.calculate_dimension(),
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "nonempty_cells": nonempty_count,
                "formula_cells": sheet_formula_count,
                "non_formula_cells": input_count,
                "comments": comment_count,
                "unlocked_cells": unlocked_count,
                "merged_ranges": [str(item) for item in ws.merged_cells.ranges],
                "hidden_rows": [index for index, dimension in ws.row_dimensions.items() if dimension.hidden],
                "hidden_columns": [index for index, dimension in ws.column_dimensions.items() if dimension.hidden],
                "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
                "auto_filter": str(ws.auto_filter.ref) if ws.auto_filter.ref else None,
                "tables": tables,
                "data_validations": validations,
                "conditional_formatting_rules": len(ws.conditional_formatting),
                "charts": len(ws._charts),
                "images": len(ws._images),
                "sheet_protected": bool(ws.protection.sheet),
            }
        )

    calculation = getattr(wb, "calculation", None)
    report = {
        "workbook": {
            "file_name": path.name,
            "format": path.suffix.lower(),
            "sheets": len(wb.sheetnames),
            "sheet_order": wb.sheetnames,
            "active_sheet": wb.active.title if wb.active else None,
            "date_system": str(getattr(wb, "epoch", "unknown")),
            "calculation_mode": getattr(calculation, "calcMode", None),
            "iterate_calculation": getattr(calculation, "iterate", None),
            "iterate_count": getattr(calculation, "iterateCount", None),
            "iterate_delta": getattr(calculation, "iterateDelta", None),
            "structure_protected": bool(getattr(wb.security, "lockStructure", False)),
            "windows_protected": bool(getattr(wb.security, "lockWindows", False)),
            "package_features": inspect_package(path),
        },
        "defined_names": defined_names(wb),
        "sheets": sheet_reports,
        "formula_summary": {
            "total_formula_cells": len(formulas),
            "function_counts": dict(function_counts.most_common()),
        },
        "formulas": formulas,
        "labels": labels,
        "values": values if include_values else None,
        "analysis_limits": [
            "Cached formula values may be absent or stale; this script does not calculate formulas.",
            "VBA presence is detected, but macro code and behavior are not interpreted.",
            "Power Query, pivots, connections, and external links require separate behavioral review.",
            "Cell-reference extraction is approximate and does not fully parse every Excel formula grammar feature.",
            "Roles and business workflow must be inferred from workbook evidence and confirmed with the user.",
        ],
    }
    wb.close()
    cached_wb.close()
    return report


def main() -> int:
    args = parse_args()
    path = args.workbook.expanduser().resolve()
    if not path.is_file():
        print(f"Workbook not found: {path}", file=sys.stderr)
        return 2
    if path.suffix.lower() not in SUPPORTED_SUFFIXES:
        print(
            f"Unsupported format {path.suffix!r}; use an OOXML workbook (.xlsx or .xlsm) or another reader.",
            file=sys.stderr,
        )
        return 2

    try:
        report = inspect_workbook(path, args.include_values, max(args.label_limit, 0))
    except (RuntimeError, zipfile.BadZipFile, OSError, ValueError) as exc:
        print(f"Could not inspect workbook: {exc}", file=sys.stderr)
        return 1

    payload = json.dumps(report, indent=2, ensure_ascii=False)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(payload + "\n", encoding="utf-8")
    else:
        print(payload)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
